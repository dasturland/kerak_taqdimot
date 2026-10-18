#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fleshkalar (USB) qanday ishlaydi mavzusidagi 10 slaydli PPTX taqdimotini yaratadi.

Talablar (foydalanuvchi tomonidan qat'iy belgilangan):
  * Aynan 10 ta slayd.
  * Sof, professional, xatosiz o'zbek tilidagi matn (verbatim).
  * Shrift faqat "Times New Roman", minimal o'lcham 12 pt.
  * 4 qism tuzilmasi: 3 ta asosiy reja + 1 ta xulosa.
  * Har bir slaydda mavzuga mos Unsplash rasmi (yuklab olinib, ichiga joylanadi).
  * Kiber / mikroelektronika / xotira qurilmalari / kvant fizikasi dizayni.

Foydalanish:
    python3 generate_presentation.py [chiqish_fayli.pptx]

Agar python-pptx o'rnatilmagan bo'lsa yoki Unsplash rasmlari yuklab
olinmasa, skript buni aniq xabar bilan bildiradi. Rasm yuklab olinmagan
holatda mavzuga mos rang-placeholder ishlatiladi (fayl baribir yaroqli
10 slaydli .pptx bo'lib qoladi).
"""
import io
import os
import sys
import urllib.request

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
except ImportError:  # python-pptx o'rnatilmagan
    sys.stderr.write(
        "XATO: python-pptx topilmadi. O'rnating: pip install python-pptx\n"
        "Internetga ulanishsiz muhit uchun 'build_pptx_fallback.py' dan "
        "foydalaning (faqat standart kutubxona kerak).\n"
    )
    raise

from slides_content import (
    SLIDES, FONT_NAME,
    COLOR_BG_DARK, COLOR_PANEL, COLOR_ACCENT, COLOR_ACCENT2,
    COLOR_TITLE, COLOR_BODY,
    SIZE_TITLE, SIZE_SUBTITLE, SIZE_SLIDE_TITLE,
    SIZE_BODY_HEAD, SIZE_BODY, SIZE_META,
)
from placeholder import make_placeholder_png

# 16:9 slayd o'lchami
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _rgb(t):
    return RGBColor(t[0], t[1], t[2])


def download_image(url, idx):
    """Unsplash rasmini yuklab oladi. Muvaffaqiyatsizlikda placeholder qaytaradi."""
    headers = {"User-Agent": "Mozilla/5.0 (presentation-generator)"}
    for attempt in range(3):
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=25) as resp:
                data = resp.read()
            if data and len(data) > 1000:
                print("  [OK] rasm yuklandi (%d bayt)" % len(data))
                return io.BytesIO(data), "jpg"
        except Exception as exc:  # noqa: BLE001
            print("  [urinish %d] rasm yuklanmadi: %s" % (attempt + 1, exc))
    print("  [FALLBACK] mavzuga mos placeholder rasm ishlatilmoqda")
    return io.BytesIO(make_placeholder_png(idx)), "png"


def set_run(run, text, size, color, bold=False, italic=False):
    run.text = text
    f = run.font
    f.name = FONT_NAME          # shriftni har bir run uchun aniq o'rnatamiz
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = _rgb(color)


def add_bg(slide, color):
    rect = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        0, 0, SLIDE_W, SLIDE_H,
    )
    rect.fill.solid()
    rect.fill.fore_color.rgb = _rgb(color)
    rect.line.fill.background()
    rect.shadow.inherit = False
    # orqa fonga joylash
    slide.shapes._spTree.remove(rect._element)
    slide.shapes._spTree.insert(2, rect._element)
    return rect


def add_accent_bar(slide, left, top, width, height, color):
    bar = slide.shapes.add_shape(1, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(color)
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_text(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def build_title_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # bo'sh layout
    add_bg(slide, COLOR_BG_DARK)

    # O'ng tomonda rasm paneli
    img_stream, _ = download_image(data["image"], 0)
    pic_left = Inches(7.4)
    slide.shapes.add_picture(img_stream, pic_left, 0, width=Inches(5.933), height=SLIDE_H)

    # Chapda aksent chiziq
    add_accent_bar(slide, Inches(0.6), Inches(1.2), Inches(0.12), Inches(2.4), COLOR_ACCENT)

    tf = add_text(slide, Inches(0.9), Inches(1.1), Inches(6.1), Inches(2.6))
    p = tf.paragraphs[0]
    set_run(p.add_run(), data["title"], SIZE_TITLE, COLOR_TITLE, bold=True)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(14)
    set_run(p2.add_run(), data["subtitle"], SIZE_SUBTITLE, COLOR_ACCENT, italic=True)

    tf2 = add_text(slide, Inches(0.9), Inches(4.2), Inches(6.1), Inches(2.6))
    first = True
    for line in data["lines"]:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph()
        p.space_after = Pt(6)
        set_run(p.add_run(), line, SIZE_META, COLOR_BODY)
        first = False
    return slide


def build_closing_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLOR_BG_DARK)

    img_stream, _ = download_image(data["image"], 9)
    slide.shapes.add_picture(img_stream, 0, 0, width=SLIDE_W, height=SLIDE_H)

    # Yarim shaffof emas, lekin to'q panel markazda
    panel = add_accent_bar(slide, Inches(1.5), Inches(2.3), Inches(10.33), Inches(3.0), COLOR_PANEL)
    panel.fill.fore_color.rgb = _rgb(COLOR_BG_DARK)

    tf = add_text(slide, Inches(1.8), Inches(2.5), Inches(9.7), Inches(2.6), MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), data["title"], SIZE_TITLE, COLOR_ACCENT, bold=True)
    for line in data["lines"]:
        pp = tf.add_paragraph()
        pp.alignment = PP_ALIGN.CENTER
        pp.space_before = Pt(8)
        set_run(pp.add_run(), line, SIZE_META, COLOR_BODY)
    return slide


def build_content_slide(prs, data, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLOR_BG_DARK)

    # Yuqori sarlavha zonasi
    add_accent_bar(slide, 0, Inches(1.15), SLIDE_W, Inches(0.05), COLOR_ACCENT)
    add_accent_bar(slide, Inches(0.6), Inches(0.35), Inches(0.12), Inches(0.6), COLOR_ACCENT)
    tf_title = add_text(slide, Inches(0.9), Inches(0.3), Inches(11.8), Inches(0.9), MSO_ANCHOR.MIDDLE)
    p = tf_title.paragraphs[0]
    set_run(p.add_run(), data["title"], SIZE_SLIDE_TITLE, COLOR_TITLE, bold=True)

    # O'ng tomonda rasm
    img_stream, _ = download_image(data["image"], idx)
    slide.shapes.add_picture(img_stream, Inches(8.5), Inches(1.6), width=Inches(4.3), height=Inches(5.4))

    # Chapda matn paneli
    panel = add_accent_bar(slide, Inches(0.5), Inches(1.6), Inches(7.7), Inches(5.4), COLOR_PANEL)
    tf = add_text(slide, Inches(0.8), Inches(1.85), Inches(7.1), Inches(4.9))
    first = True
    for line in data["lines"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.space_after = Pt(10)
        bold = line.strip().startswith("-")
        set_run(p.add_run(), line, SIZE_BODY if bold else SIZE_BODY_HEAD,
                COLOR_BODY if bold else COLOR_ACCENT, bold=not bold)
        first = False
    return slide


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "fleshka_usb_taqdimot.pptx"
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for idx, data in enumerate(SLIDES):
        print("Slayd %d/%d: %s" % (idx + 1, len(SLIDES), data["title"]))
        if data["kind"] == "title":
            build_title_slide(prs, data)
        elif data["kind"] == "closing":
            build_closing_slide(prs, data)
        else:
            build_content_slide(prs, data, idx)

    assert len(prs.slides._sldIdLst) == 10, "Slaydlar soni 10 emas!"
    prs.save(out)
    print("\nTayyor: %s (slaydlar soni: %d)" % (os.path.abspath(out), len(prs.slides._sldIdLst)))


if __name__ == "__main__":
    main()
